"""1688 promotion equals SUMIFS(spend, product)/COUNTIFS(original rows)."""
from io import BytesIO
import re,zipfile
import polars as pl
import pytest
from conftest import MODELS,write_xlsx
from ledger.engine.runtime import Ingestion,run,ingest
from ledger.engine.parse import _read_with_openpyxl
from ledger.model.loader import load_model
from ledger.model.schema import Model,Store,SourceContract,Metric,StatementNode,ParseOptions,Template,ColumnBinding
from test_promotion_dates import ingested
from ledger import view

HEADER=['商品ID','offer标题','曝光量','点击量','消耗量','日期']


def scenario(promo_rows,*,live=True,dated=True):
    built=load_model(MODELS/'cn-ecommerce')
    model=Model(id='1688-test',name='test',stores=(Store(id='s',name='shop',platform='alibaba1688'),),
        sources=(SourceContract(id='order_detail',name='订单',is_spine=True,owner_role='operations',cadence='monthly'),
                 SourceContract(id='promotion',name='推广',owner_role='operations',cadence='monthly',required_for_close=False)),
        metrics=(built.metric('ad_cost'),Metric(id='sales',name='订单金额',source='order_detail',value={'op':'sum','of':['buyer_paid']})),
        statement=(StatementNode(id='ad',name='推广费用',formula={'op':'add','of':['ad_cost']}),))
    headers=['订单号','买家公司名称','货品总价','实付款（元）','运单号','订单创建时间','Offer_ID','数量']
    rows=[['A','buyer',10,10,'T','2026-06-30','P1',1],
          ['B','buyer',10,10,'T','2026-07-01','P1',1],['C','buyer',90,90,'T','2026-07-02','P1',9],
          ['D','buyer',10,10,'T','2026-07-03','P2',1],['E','buyer',10,10,'T','2026-07-04','P3',1]]
    original=ingested(built.template('alibaba1688_order_detail_v1'),headers,rows)
    original.frame=original.frame.with_columns(pl.lit('shop').alias('__hint_store__'))
    promo=ingested(built.template('promotion_1688_v1'),HEADER if dated else HEADER[:-1],promo_rows)
    promo.frame=promo.frame.with_columns(pl.lit('shop').alias('__hint_store__'),pl.lit('2026-07').alias('__hint_period__'))
    items=[original,promo]
    if live:
        fields=['order_id','product_id','store_name','order_time','buyer_paid']
        template=Template(id='order_console_test',source='order_detail',match_columns=('order_id',),
            time_slots={'order_date':'order_time'},bindings=tuple(ColumnBinding(role=f,columns=(f,)) for f in fields))
        items.append(ingested(template,fields,[['EXTRA','P1','shop','2026-07-05',999]]))
    return run(Ingestion(model=model,items=items),'alibaba1688')


@pytest.mark.parametrize('live',[False,True])
def test_sumifs_countifs_uses_original_rows_not_quantity_or_live_additions(live):
    result=scenario([['P1','a',1,1,12,'2026-07-01'],['P1','a',1,1,18,'2026-07-02'],
                     ['P2','b',1,1,8,'2026-07-03'],['合计','合计',1,1,38,'总计']],live=live)
    sl=result.slice('shop','2026-07')
    assert sl.nodes['ad'].value==-38 and sl.nodes['ad'].available
    posted=result.spine_facts.filter(pl.col('metric_id')=='ad_cost')
    assert sorted(posted['amount'].to_list())==[-15.,-15.,-8.]
    assert posted['period'].unique().to_list()==['2026-07']
    assert result.facts.filter(pl.col('link_key')=='__store_wide__').is_empty()
    assert sl.facts.filter(pl.col('metric_id')=='ad_cost')['contribution'].sum()==-38


def test_monthly_file_without_date_uses_explicit_period():
    result=scenario([['P1','a',1,1,'30']],dated=False)
    assert result.slice('shop','2026-07').nodes['ad'].value==-30
    assert result.facts.filter(pl.col('metric_id')=='ad_cost')['source_period'].unique().to_list()==['2026-07']


def test_unmatched_spend_is_not_automatically_store_wide_or_hidden_by_footer():
    result=scenario([['P1','a',1,1,30,'2026-07-01'],['UNKNOWN','x',1,1,7,'2026-07-01'],['合计','合计',1,1,37,'总计']])
    sl=result.slice('shop','2026-07')
    row=view._statement(sl,result.model)[0]
    assert row['verified_partial']==-30 and not row['available']
    orphan=sl.facts.filter(pl.col('link_key')=='UNKNOWN')
    assert orphan['contribution'].item()==0 and orphan['booking_status'].item()=='unposted'


def test_promotion_template_recognizes_actual_header_and_numeric_spend(tmp_path):
    model=load_model(MODELS/'cn-ecommerce')
    path=write_xlsx(tmp_path/'推广-1688-2607月.xlsx',[HEADER[:-1],['681267612667','标题',10,2,'377.17']])
    parsed=ingest([path],model,default_store='shop')
    assert len(parsed.items)==1 and parsed.items[0].ok
    assert parsed.items[0].template.id=='promotion_1688_v1'
    assert parsed.items[0].frame['spend'].item()==377.17


def test_openpyxl_fallback_reads_bad_a1_dimension(tmp_path):
    path=write_xlsx(tmp_path/'promotion.xlsx',[HEADER[:-1],['681267612667','标题',10,2,377.17],['2','第二个',1,1,1]])
    original=path.read_bytes();buf=BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as src,zipfile.ZipFile(buf,'w') as dst:
        for name in src.namelist():
            data=src.read(name)
            if name=='xl/worksheets/sheet1.xml':data=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1"',data)
            dst.writestr(name,data)
    # A bytes-backed input avoids relying on the filename's declared rectangle.
    sheets=_read_with_openpyxl(tmp_path/'raw.bin',buf.getvalue(),ParseOptions())
    assert len(sheets[0][1])==3 and sheets[0][1][0]==HEADER[:-1]
    assert sheets[0][1][1][-1]==377.17


def test_merged_order_dates_do_not_shrink_countifs_denominator(tmp_path):
    import openpyxl
    model=load_model(MODELS/'cn-ecommerce');store=model.store('alibaba1688_mt2r23jf')
    headers=['订单号','买家公司名称','货品总价','实付款（元）','运单号','订单创建时间','Offer_ID','数量']
    orders=write_xlsx(tmp_path/'订单明细-2607月.xlsx',[headers,
        ['A','buyer',100,100,'T','2026-07-01','P1',1],['A','buyer',100,100,'T','2026-07-01','P1',9]])
    wb=openpyxl.load_workbook(orders);sh=wb.active
    sh.merge_cells('A2:A3');sh.merge_cells('F2:F3');wb.save(orders);wb.close()
    promotion=write_xlsx(tmp_path/'推广-2607月.xlsx',[HEADER[:-1],['P1','商品',100,20,30]])
    result=run(ingest([orders,promotion],model,default_store=store.name),'alibaba1688')
    posted=result.spine_facts.filter(pl.col('metric_id')=='ad_cost')
    assert posted.height==2 and posted['amount'].to_list()==[-15.,-15.]
