import io
import json
from datetime import datetime

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from ledger.commission_registry import Registry, RegistryError, RevisionConflict
from ledger.commission_batch import preview, apply, parse_excel
from ledger.commission_catalog import settings
from ledger.commission_api import install
from ledger.model.schema import Store
from ledger.workspace import Workspace
from test_commission import _model


def setup(tmp_path):
    model=_model(stores=(Store(id='s1',name='甲店',platform='taobao'),Store(id='s2',name='乙店',platform='taobao')))
    r=Registry(tmp_path)
    a=r.person_save({'name':'甲'},'test','登记');b=r.person_save({'name':'乙'},'test','登记')
    for sid in ['s1','s2']:
        r.save_setting({'store_id':sid,'product_id':'123456789001','valid_from':'2000-01-01',
                        'allocations':[{'person_id':a['id'],'rate':'.03'},{'person_id':b['id'],'rate':'.02'}]},'test')
    return r,model,a,b


def request():
    return {'targets':[{'store_id':sid,'product_id':'123456789001','revision':1} for sid in ['s1','s2']],
            'template':{'valid_from':'2001-01-01','mode':'distribute','allocations':[{'name':'新增丙','rate':'.01'}]},'operation':'merge'}


def test_cross_store_preview_no_changes_apply_once_and_keep_others(tmp_path):
    r,m,a,b=setup(tmp_path);before=r.revision()
    plan=preview(r,m,request(),'test')
    assert r.revision()==before and len(r.people())==2
    assert plan['stores']==2 and plan['count']==2 and len(plan['rows'][0]['after'])==3
    result=apply(r,m,plan['id'],'test');rev=r.revision()
    assert result['count']==2 and len(r.people())==3
    assert apply(r,m,plan['id'],'test')==result and r.revision()==rev
    assert len(settings(r,person_id=a['id'])['rows'])==2
    for row in settings(r)['rows']:
        assert sum(float(p['rate']) for p in row['people'])==pytest.approx(.06)
    assert r.active('s1')[1][0]['body']['segments'][0]['valid_to']=='2001-01-01T00:00:00'


def test_merge_retains_existing_duty_and_hierarchy_source(tmp_path):
    r,m,a,b=setup(tmp_path)
    r.save_setting({'store_id':'s1','product_id':'123456789001','expected_revision':1,'valid_from':'2000-06-01',
                    'allocations':[{'person_id':a['id'],'rate':'.03','duty':'produce'},
                                   {'person_id':b['id'],'rate':'.02','duty':'cut','source':'hierarchy'}]},'test')
    body={'targets':[{'store_id':'s1','product_id':'123456789001','revision':2}],
          'template':{'valid_from':'2001-01-01','mode':'distribute','allocations':[{'person_id':a['id'],'rate':'.04','duty':'produce'}]},'operation':'merge'}
    plan=preview(r,m,body,'test')
    other=next(p for p in plan['rows'][0]['after'] if p['person_id']==b['id'])
    assert other['duty']=='cut' and other['source']=='hierarchy'
    apply(r,m,plan['id'],'test')
    actual=r.active('s1')[1][0]['body']['segments'][-1]['allocations']
    other=next(p for p in actual if p['person_id']==b['id'])
    assert other['duty']=='cut' and other['source']=='hierarchy'


def test_conflict_rolls_back_every_target_and_new_person(tmp_path):
    r,m,a,b=setup(tmp_path);plan=preview(r,m,request(),'test')
    r.save_setting({'store_id':'s2','product_id':'123456789001','expected_revision':1,'valid_from':'2000-06-01',
                    'allocations':[{'person_id':a['id'],'rate':'.05'}]},'other')
    before=r.revision()
    with pytest.raises(RevisionConflict):apply(r,m,plan['id'],'test')
    assert r.revision()==before and len(r.people())==2
    assert r.active('s1')[1][0]['revision']==1


def test_filtered_bulk_remove_and_roster_counts(tmp_path):
    r,m,a,b=setup(tmp_path)
    plan=preview(r,m,{'scope':{'store_id':'s1','person_id':a['id']},'operation':'remove',
                     'template':{'valid_from':'2001-01-01','allocations':[{'person_id':a['id']}]}},'test')
    apply(r,m,plan['id'],'test')
    assert len(settings(r,person_id=a['id'])['rows'])==1
    ws=Workspace(tmp_path);app=FastAPI();install(app,lambda:ws,lambda:m);c=TestClient(app)
    people=c.get('/api/commission-v2/people/summary').json()['people']
    assert next(p for p in people if p['id']==a['id'])['products']==1
    assert next(p for p in people if p['id']==b['id'])['stores']==2
    assert c.get('/api/commission-v2/settings?store_id=s2').json()['rows'][0]['store_id']=='s2'


def excel(rows, second=None):
    wb=openpyxl.Workbook();s=wb.active;s.title='提成设置'
    columns=['店铺','宝贝ID','商品名称','人员','提成比例','状态','生效时间','结束时间']
    s.append(columns)
    for row in rows:s.append(row)
    if second:
        t=wb.create_sheet('第二家店');t.append(columns)
        for row in second:t.append(row)
    wb.create_sheet('填写说明').append(['不要把示例自动导入'])
    f=io.BytesIO();wb.save(f);return f.getvalue()


def test_excel_multi_sheet_multi_person_long_ids_and_dates(tmp_path):
    r,m,a,b=setup(tmp_path)
    raw=excel([['甲店','123456789012345678','新宝贝','甲','3%','提成中',datetime(2001,1,1),None],
               ['甲店','123456789012345678','新宝贝','乙','2%','提成中',datetime(2001,1,1),None]],
              [['s2','123456789002','第二宝贝','新增丁',6.5,'提成中','2001-01-01','']])
    parsed=parse_excel(raw,'test.xlsx',m)
    assert len(parsed['changes'])==2 and parsed['rows']==3
    assert parsed['changes'][0]['product_id']=='123456789012345678'
    plan=preview(r,m,parsed,'test');assert plan['new_count']==2 and len(r.people())==2
    apply(r,m,plan['id'],'test')
    assert len(r.people())==3
    assert len(settings(r)['rows'])==4


def test_excel_reports_ambiguous_rates_formulas_and_numeric_long_ids(tmp_path):
    _,m,_,_=setup(tmp_path)
    raw=excel([['甲店','123456789001','宝贝','甲',.05,'提成中','2001-01-01',''],
               ['甲店',123456789012345678,'宝贝','甲','5%','提成中','2001-01-01',''],
               ['甲店','123456789003','宝贝','甲','=1/20','提成中','2001-01-01','']])
    parsed=parse_excel(raw,'bad.xlsx',m)
    assert len(parsed['errors'])==3
    assert '小数比例' in parsed['errors'][0]['error']
    assert '长宝贝ID' in parsed['errors'][1]['error']
    assert '公式' in parsed['errors'][2]['error']


def test_http_import_preview_and_commit_without_business_side_effects_before_confirm(tmp_path):
    r,m,a,b=setup(tmp_path);ws=Workspace(tmp_path);app=FastAPI();install(app,lambda:ws,lambda:m);c=TestClient(app)
    raw=excel([['甲店','123456789002','新宝贝','甲','5%','提成中','2001-01-01','']])
    before=r.revision();p=c.post('/api/commission-v2/settings/import-preview',files={'file':('new.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert p.status_code==200 and r.revision()==before
    bid=p.json()['id'];assert c.post('/api/commission-v2/settings/apply/'+bid).json()['applied']
    assert c.post('/api/commission-v2/settings/apply/'+bid).json()['count']==1


def test_select_all_honors_exclusions_and_ignores_unmapped_stores(tmp_path):
    r,m,a,b=setup(tmp_path)
    with r.transaction() as c:
        c.execute("INSERT INTO catalog VALUES(?,?,?,?,?,?)",('unmapped:9','123456789999','未映射商品','9','{}','2000-01-01'))
    plan=preview(r,m,{'scope':{'excluded':['s1:123456789001']},'template':{'valid_from':'2001-01-01','mode':'exclude'}},'test')
    assert plan['count']==1 and plan['rows'][0]['store_id']=='s2'
    apply(r,m,plan['id'],'test')
    assert settings(r,store_id='s1')['rows'][0]['state']=='enabled'
    assert settings(r,store_id='s2')['rows'][0]['state']=='disabled'


def test_person_filter_includes_scheduled_assignment_after_gap(tmp_path):
    r,m,a,b=setup(tmp_path)
    r.save_scheme('s1','123456789005',{'segments':[
        {'valid_from':'2000-01-01','valid_to':'2001-01-01','mode':'exclude','allocations':[],'total_rate':'0'},
        {'valid_from':'2999-01-01','mode':'distribute','total_rate':'.05','allocations':[{'person_id':a['id'],'rate':'.05'}]}]},'test','后续安排',publish=True)
    row=next(x for x in settings(r,person_id=a['id'],at='2026-09-01T00:00:00')['rows'] if x['product_id']=='123456789005')
    assert row['state']=='scheduled' and row['setting']['valid_from']=='2999-01-01T00:00:00'


def test_remove_unassigned_person_does_not_create_empty_settings(tmp_path):
    r,m,a,b=setup(tmp_path);revision=r.revision()
    with pytest.raises(RegistryError,match='没有需要调整'):
        preview(r,m,{'targets':[{'store_id':'s1','product_id':'123456789003'}],
                     'operation':'remove','template':{'valid_from':'2001-01-01','allocations':[{'person_id':a['id']}]}},'test')
    assert r.revision()==revision and len(settings(r)['rows'])==2


def test_bulk_replace_without_end_overwrites_scheduled_future_but_single_edit_keeps_it(tmp_path):
    r,m,a,b=setup(tmp_path)
    r.save_setting({'store_id':'s1','product_id':'123456789001','expected_revision':1,
                    'valid_from':'2002-01-01','allocations':[{'person_id':a['id'],'rate':'.05'}]},'test')
    plan=preview(r,m,{'targets':[{'store_id':'s1','product_id':'123456789001','revision':2}],
        'operation':'replace','template':{'valid_from':'2001-01-01','mode':'distribute',
        'allocations':[{'person_id':b['id'],'rate':'.05'}]}},'test')
    assert plan['future_overwritten_count']==1
    assert plan['rows'][0]['valid_to']=='' and plan['rows'][0]['future_overwritten']==1
    apply(r,m,plan['id'],'test')
    segments=r.active('s1')[1][0]['body']['segments']
    assert [(x['valid_from'],x['valid_to']) for x in segments]==[
        ('2000-01-01T00:00:00','2001-01-01T00:00:00'),('2001-01-01T00:00:00','')]
    assert segments[-1]['allocations'][0]['person_id']==b['id']

    # The ordinary one-item editor still protects an explicitly scheduled change.
    r.save_setting({'store_id':'s2','product_id':'123456789001','expected_revision':1,
                    'valid_from':'2002-01-01','allocations':[{'person_id':a['id'],'rate':'.05'}]},'test')
    kept=r.save_setting({'store_id':'s2','product_id':'123456789001','expected_revision':2,
                         'valid_from':'2001-01-01','allocations':[{'person_id':b['id'],'rate':'.05'}]},'test')
    assert [(x['valid_from'],x['valid_to']) for x in kept['body']['segments']]==[
        ('2000-01-01T00:00:00','2001-01-01T00:00:00'),
        ('2001-01-01T00:00:00','2002-01-01T00:00:00'),('2002-01-01T00:00:00','')]
